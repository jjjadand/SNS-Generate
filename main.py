from fastapi import FastAPI, HTTPException, Depends, UploadFile, File, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel
from typing import Any, Dict, List, Optional
import sqlite3
import json
import os
import tomllib
import httpx
import base64
import shutil
import re
from pathlib import Path
from datetime import datetime
from openai import OpenAI

BASE_DIR = Path(__file__).resolve().parent

try:
    from dotenv import load_dotenv
    load_dotenv(BASE_DIR / ".env")
except ImportError:
    pass

# Optional imports
try:
    from docx import Document as DocxDocument
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    from bs4 import BeautifulSoup
    HAS_BS4 = True
except ImportError:
    HAS_BS4 = False

UPLOAD_DIR = BASE_DIR / "uploads"
UPLOAD_DIR.mkdir(exist_ok=True)

# ── Load Codex config & auth ───────────────────────────────────
def load_codex_config() -> dict:
    """Read base_url and api_key from local Codex config (~/.codex/)"""
    result = {"base_url": None, "api_key": None, "model": "gpt-4o"}

    # Read config.toml for base_url and model
    config_path = Path.home() / ".codex" / "config.toml"
    if config_path.exists():
        try:
            with open(config_path, "rb") as f:
                cfg = tomllib.load(f)
            provider_name = cfg.get("model_provider")
            result["model"] = cfg.get("model", "gpt-4o")
            if provider_name and "model_providers" in cfg:
                provider = cfg["model_providers"].get(provider_name, {})
                result["base_url"] = provider.get("base_url")
        except Exception as e:
            print(f"[warn] Failed to read codex config: {e}")

    # Read auth.json for api_key
    auth_path = Path.home() / ".codex" / "auth.json"
    if auth_path.exists():
        try:
            with open(auth_path) as f:
                auth = json.load(f)
            result["api_key"] = auth.get("OPENAI_API_KEY")
        except Exception as e:
            print(f"[warn] Failed to read codex auth: {e}")

    return result

CODEX_CONFIG = load_codex_config()

# Fix base_url: ensure it ends with /v1
if CODEX_CONFIG.get("base_url"):
    bu = CODEX_CONFIG["base_url"].rstrip("/")
    if not bu.endswith("/v1"):
        bu = bu + "/v1"
    CODEX_CONFIG["base_url"] = bu

# Remove ALL_PROXY from env (socks:// breaks httpx used by openai SDK)
for _k in ("ALL_PROXY", "all_proxy"):
    os.environ.pop(_k, None)

print(f"[init] Codex provider base_url: {CODEX_CONFIG['base_url']}")
print(f"[init] Codex model: {CODEX_CONFIG['model']}")
print(f"[init] Codex api_key loaded: {'yes' if CODEX_CONFIG['api_key'] else 'no'}")

app = FastAPI(title="SNS Marketing Hub API", version="1.0.0")
FRONTEND_PATH = Path(__file__).with_name("sns-marketing-hub.html")

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "null",
        "http://localhost:5173",
        "http://localhost:4173",
        "http://127.0.0.1:5173",
    ],
    allow_origin_regex=r"https?://(localhost|127\.0\.0\.1)(:\d+)?",
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# ── Database Setup ─────────────────────────────────────────────
DB_PATH = BASE_DIR / "sns_data.db"

def get_db():
    conn = sqlite3.connect(DB_PATH, check_same_thread=False)
    conn.row_factory = sqlite3.Row
    try:
        yield conn
    finally:
        conn.close()

def init_db():
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()

    c.execute("""
        CREATE TABLE IF NOT EXISTS materials (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            type TEXT NOT NULL,
            content TEXT,
            file_path TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)

    c.execute("""
        CREATE TABLE IF NOT EXISTS crawl_sources (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            kind TEXT NOT NULL,
            name TEXT NOT NULL,
            platform TEXT DEFAULT '',
            region TEXT DEFAULT '',
            url TEXT NOT NULL UNIQUE,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)

    c.execute("""
        CREATE TABLE IF NOT EXISTS live_monthly_snapshots (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            source_id INTEGER NOT NULL,
            month TEXT NOT NULL,
            year INTEGER NOT NULL,
            linkedin_impressions INTEGER DEFAULT 0,
            x_impressions INTEGER DEFAULT 0,
            linkedin_engagement REAL DEFAULT 0,
            x_engagement REAL DEFAULT 0,
            linkedin_followers INTEGER DEFAULT 0,
            x_followers INTEGER DEFAULT 0,
            linkedin_clicks INTEGER DEFAULT 0,
            x_clicks INTEGER DEFAULT 0,
            raw_excerpt TEXT DEFAULT '',
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)

    c.execute("""
        CREATE TABLE IF NOT EXISTS live_posts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            source_id INTEGER NOT NULL,
            platform TEXT NOT NULL,
            title TEXT NOT NULL,
            content TEXT DEFAULT '',
            post_type TEXT DEFAULT '',
            impressions INTEGER DEFAULT 0,
            likes INTEGER DEFAULT 0,
            comments INTEGER DEFAULT 0,
            shares INTEGER DEFAULT 0,
            engagement_rate REAL DEFAULT 0,
            posted_at TEXT DEFAULT '',
            month INTEGER,
            year INTEGER,
            source_url TEXT DEFAULT '',
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)

    c.execute("""
        CREATE TABLE IF NOT EXISTS competitor_snapshots (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            source_id INTEGER NOT NULL,
            name TEXT NOT NULL,
            platform TEXT DEFAULT '',
            region TEXT DEFAULT '',
            followers TEXT DEFAULT '',
            posting_freq TEXT DEFAULT '',
            avg_engagement TEXT DEFAULT '',
            top_actions_json TEXT DEFAULT '[]',
            learnings TEXT DEFAULT '',
            raw_excerpt TEXT DEFAULT '',
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)

    c.execute("""
        CREATE TABLE IF NOT EXISTS sns_monthly (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            month TEXT NOT NULL,
            year INTEGER NOT NULL,
            linkedin_impressions INTEGER DEFAULT 0,
            x_impressions INTEGER DEFAULT 0,
            linkedin_engagement REAL DEFAULT 0,
            x_engagement REAL DEFAULT 0,
            linkedin_followers INTEGER DEFAULT 0,
            x_followers INTEGER DEFAULT 0,
            linkedin_clicks INTEGER DEFAULT 0,
            x_clicks INTEGER DEFAULT 0,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(month, year)
        )
    """)

    c.execute("""
        CREATE TABLE IF NOT EXISTS posts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            platform TEXT NOT NULL,
            title TEXT NOT NULL,
            content TEXT,
            post_type TEXT,
            impressions INTEGER DEFAULT 0,
            likes INTEGER DEFAULT 0,
            comments INTEGER DEFAULT 0,
            shares INTEGER DEFAULT 0,
            engagement_rate REAL DEFAULT 0,
            posted_at TEXT,
            month INTEGER,
            year INTEGER,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)

    c.execute("""
        CREATE TABLE IF NOT EXISTS kols (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            handle TEXT,
            platform TEXT,
            region TEXT,
            category TEXT,
            followers TEXT,
            avg_views TEXT,
            engagement_rate TEXT,
            email TEXT,
            match_score INTEGER DEFAULT 0,
            status TEXT DEFAULT '待联系',
            language TEXT DEFAULT 'English',
            notes TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)

    c.execute("""
        CREATE TABLE IF NOT EXISTS market_reports (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL,
            industry TEXT DEFAULT '',
            region TEXT DEFAULT 'Global',
            language TEXT DEFAULT 'zh-CN',
            report_content TEXT DEFAULT '',
            market_data TEXT DEFAULT '',
            competitors TEXT DEFAULT '',
            trends TEXT DEFAULT '',
            generated_by TEXT DEFAULT 'local-codex',
            created_at TEXT DEFAULT CURRENT_TIMESTAMP,
            updated_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)

    c.execute("""
        CREATE TABLE IF NOT EXISTS market_news_cache (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            keyword TEXT NOT NULL,
            title TEXT NOT NULL,
            snippet TEXT DEFAULT '',
            url TEXT DEFAULT '',
            source TEXT DEFAULT '',
            published_at TEXT DEFAULT '',
            fetched_at TEXT DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(keyword, title)
        )
    """)

    conn.commit()
    conn.close()

init_db()


@app.get("/", include_in_schema=False)
def index():
    if not FRONTEND_PATH.exists():
        raise HTTPException(status_code=404, detail="Frontend page not found")
    return FileResponse(
        FRONTEND_PATH,
        headers={
            "Cache-Control": "no-store, no-cache, must-revalidate, max-age=0",
            "Pragma": "no-cache",
            "Expires": "0",
        },
    )


@app.get("/sns-marketing-hub.html", include_in_schema=False)
def frontend_html():
    if not FRONTEND_PATH.exists():
        raise HTTPException(status_code=404, detail="Frontend page not found")
    return FileResponse(
        FRONTEND_PATH,
        headers={
            "Cache-Control": "no-store, no-cache, must-revalidate, max-age=0",
            "Pragma": "no-cache",
            "Expires": "0",
        },
    )

# ── Pydantic Models ────────────────────────────────────────────
class GenerateRequest(BaseModel):
    content_type: str       # linkedin, edm, landing, blog, seo, case, wechat
    platform: Optional[str] = "LinkedIn"
    post_type: Optional[str] = "product"
    topic: Optional[str] = ""
    product: Optional[str] = ""
    features: Optional[str] = ""
    audience: Optional[str] = ""
    highlights: Optional[str] = ""
    cta_url: Optional[str] = ""
    subject_line: Optional[str] = ""
    keyword: Optional[str] = ""
    region: Optional[str] = "Global"
    customer: Optional[str] = ""
    industry: Optional[str] = ""
    pain_point: Optional[str] = ""
    solution: Optional[str] = ""
    blog_length: Optional[str] = "medium"
    angle: Optional[str] = "product"
    api_key: Optional[str] = ""
    material_ids: Optional[List[int]] = []   # ← 新增：引用素材库 ID
    platforms: Optional[List[str]] = []
    language: Optional[str] = "zh-CN"
    custom_prompt: Optional[str] = ""

class URLIngestRequest(BaseModel):
    url: str
    name: Optional[str] = ""


class CrawlSourceIn(BaseModel):
    kind: str            # monthly / post / competitor
    name: str
    url: str
    platform: Optional[str] = ""
    region: Optional[str] = ""


class CrawlRefreshRequest(BaseModel):
    kind: Optional[str] = None
    source_ids: Optional[List[int]] = []

class SNSMonthlyData(BaseModel):
    month: str
    year: int
    linkedin_impressions: int = 0
    x_impressions: int = 0
    linkedin_engagement: float = 0
    x_engagement: float = 0
    linkedin_followers: int = 0
    x_followers: int = 0
    linkedin_clicks: int = 0
    x_clicks: int = 0

class PostData(BaseModel):
    platform: str
    title: str
    content: Optional[str] = ""
    post_type: Optional[str] = ""
    impressions: int = 0
    likes: int = 0
    comments: int = 0
    shares: int = 0
    engagement_rate: float = 0
    posted_at: Optional[str] = ""
    month: Optional[int] = None
    year: Optional[int] = None

class KOLData(BaseModel):
    name: str
    handle: Optional[str] = ""
    platform: Optional[str] = ""
    region: Optional[str] = ""
    category: Optional[str] = ""
    followers: Optional[str] = ""
    avg_views: Optional[str] = ""
    engagement_rate: Optional[str] = ""
    email: Optional[str] = ""
    match_score: Optional[int] = 0
    status: Optional[str] = "待联系"
    language: Optional[str] = "English"
    notes: Optional[str] = ""

class KOLStatusUpdate(BaseModel):
    status: str


class MarketReportRequest(BaseModel):
    industry: Optional[str] = ""
    region: Optional[str] = "Global"
    keywords: Optional[List[str]] = []
    language: Optional[str] = "zh-CN"
    report_type: Optional[str] = "comprehensive"   # brief / comprehensive / competitor


class MarketReportData(BaseModel):
    title: str
    industry: Optional[str] = ""
    region: Optional[str] = "Global"
    language: Optional[str] = "zh-CN"
    report_content: Optional[str] = ""
    market_data: Optional[str] = ""
    competitors: Optional[str] = ""
    trends: Optional[str] = ""

# ── AI Generate Endpoint ───────────────────────────────────────
SOCIAL_PLATFORM_GUIDES_ZH = {
    "微信公众号": "输出顺序：1. 标题 2. 导语 3. 正文（带2-4个小标题） 4. 结尾互动引导。篇幅 800-1200 字，适合国内 B2B/科技受众阅读。",
    "小红书": "输出顺序：1. 标题 2. 正文 3. 话题标签。语气真诚、种草感强、信息密度高，适合图文笔记。",
    "抖音": "输出顺序：1. 标题 2. 30-60 秒口播脚本 3. 分镜/字幕提示 4. 结尾 CTA。节奏快、金句前置。",
    "视频号": "输出顺序：1. 标题 2. 发布文案 3. 30-60 秒口播脚本。适合企业账号发布。",
    "B站": "输出顺序：1. 标题 2. 简介 3. 视频大纲。强调技术点、适合工程师观看。",
    "知乎": "输出顺序：1. 标题 2. 摘要 3. 正文。结构清晰，强调专业分析和可验证信息。",
    "LinkedIn": "严格使用结构：第一行是一句话标题；第二段用一句话做整体介绍；然后分点介绍核心卖点、实现步骤或参与亮点；最后一行写 CTA，并放入指定链接。语气克制、可信，适合企业和技术决策者。",
    "X": "更短、更直接。用 1 句 hook + 1-2 个关键信息点 + CTA 链接；尽量控制在单条短帖长度内，避免长段落和复杂铺垫。",
}

SOCIAL_PLATFORM_GUIDES_EN = {
    "LinkedIn": "Use this exact structure: one-line headline, one-sentence overview, bullet points for value points or implementation steps, and a final CTA with the provided link.",
    "X": "Keep it much shorter: one hook, one or two key points, and a CTA link. Avoid long paragraphs.",
    "微信公众号": "Output in Chinese even if other fields are English. Use a long-form article structure.",
}

CONTENT_CATEGORY_GUIDES_ZH = {
    "demo_wiki": "Demo / Wiki 教程介绍：目标是让用户点击教程并复现。重点写清楚用户能学到什么、适合谁、关键实现步骤、需要的硬件/软件、最终能跑出什么效果。",
    "event": "活动宣传：目标是让用户报名、预约或观看。重点写清楚活动主题、时间/地点/形式、适合人群、议程亮点、参与收益和报名/观看 CTA。",
    "product_intro": "产品介绍：目标是让用户了解产品价值并点击了解/购买/咨询。重点写清楚使用场景、核心卖点、规格/生态优势、部署价值和 CTA。",
}

CONTENT_CATEGORY_GUIDES_EN = {
    "demo_wiki": "Demo / Wiki tutorial: drive users to open the tutorial and reproduce the demo. Explain what they will learn, who it is for, key implementation steps, required hardware/software, and expected result.",
    "event": "Event promotion: drive registrations, appointments, or attendance. Cover the theme, time/place/format, audience, agenda highlights, user benefit, and registration/viewing CTA.",
    "product_intro": "Product introduction: drive product understanding and clicks for details, purchase, or consultation. Cover use cases, core value points, specs/ecosystem strengths, deployment value, and CTA.",
}

CONTENT_CATEGORY_ALIASES = {
    "demo": "demo_wiki",
    "wiki": "demo_wiki",
    "tutorial": "demo_wiki",
    "demo/wiki教程介绍": "demo_wiki",
    "demo/wiki 教程介绍": "demo_wiki",
    "教程介绍": "demo_wiki",
    "活动宣传": "event",
    "活动预热": "event",
    "event": "event",
    "产品介绍": "product_intro",
    "产品发布": "product_intro",
    "product": "product_intro",
    "product_intro": "product_intro",
}


def is_chinese_request(req: GenerateRequest) -> bool:
    language = (req.language or "").lower()
    platform = req.platform or ""
    if "zh" in language or "cn" in language:
        return True
    return any(token in platform for token in ("微信", "小红书", "抖音", "视频号", "B站", "知乎", "中文"))


def normalize_content_category(post_type: Optional[str]) -> str:
    raw = (post_type or "product_intro").strip()
    return CONTENT_CATEGORY_ALIASES.get(raw, raw if raw in CONTENT_CATEGORY_GUIDES_ZH else "product_intro")


def build_social_prompt(req: GenerateRequest) -> str:
    topic = req.topic or req.product or "未命名主题"
    audience = req.audience or "国内科技行业用户、工程师、方案负责人"
    features = req.features or req.highlights or "未提供"
    platform = req.platform or "微信公众号"
    cta_url = (req.cta_url or "").strip() or "未提供，请在 CTA 位置保留 [CTA链接]"
    category = normalize_content_category(req.post_type)
    is_zh = is_chinese_request(req)
    guide_map = SOCIAL_PLATFORM_GUIDES_ZH if is_zh else SOCIAL_PLATFORM_GUIDES_EN
    platform_guide = guide_map.get(platform, guide_map["微信公众号" if is_zh else "LinkedIn"])
    category_guide = (CONTENT_CATEGORY_GUIDES_ZH if is_zh else CONTENT_CATEGORY_GUIDES_EN)[category]

    if is_zh:
        return f"""你是 Seeed Studio（矽递科技）的资深中文内容策划，负责为不同平台生成适配内容。

目标平台：{platform}
主题 / 产品：{topic}
内容类型：{category}
目标受众：{audience}
关键信息：{features}
CTA 链接：{cta_url}

平台要求：
{platform_guide}

内容类型要求：
{category_guide}

统一要求：
- 默认使用中文输出，必要时保留产品型号、英文术语
- 不能空泛，要尽量具体、可执行、像真实市场内容
- 根据内容类型优先突出教程价值、活动参与价值或产品卖点
- 语气自然，不要明显 AI 腔
- CTA 必须放在结尾；如果提供了 CTA 链接，必须原样包含该链接
- 不要解释你的思路，只输出最终可直接发布/再编辑的内容
"""

    return f"""You are a senior content strategist for Seeed Studio.

Target platform: {platform}
Topic / Product: {topic}
Content category: {category}
Target audience: {audience}
Key details: {features}
CTA link: {cta_url}

Platform requirements:
{platform_guide}

Content category requirements:
{category_guide}

Global requirements:
- Be specific and commercially useful
- Match the content category: tutorial value, event participation value, or product value
- Put the CTA at the end; if a CTA link is provided, include it exactly
- Avoid generic AI-sounding filler
- Output only the final ready-to-edit content
"""


def build_prompt(req: GenerateRequest) -> str:
    if req.content_type == "social":
        return build_social_prompt(req)

    prompts = {
        "linkedin": f"""You are an expert B2B tech copywriter for Seeed Studio, a global edge AI hardware company.

Write a compelling {req.platform} post about: {req.topic or req.product}
Post type: {req.post_type} (product_launch / industry_insight / case_study)

Guidelines:
- LinkedIn: 150-300 words, professional tone, 3-5 relevant hashtags, use emojis sparingly but strategically
- X (Twitter): Max 280 chars for main post, punchy and opinionated, 2-3 hashtags
- Focus on value for engineers, developers, and tech decision-makers
- Avoid generic corporate speak
- End with a clear call-to-action

Write ONLY the post content, no explanations.""",

        "edm": f"""You are an expert email marketer for Seeed Studio.

Write a weekly EDM (email marketing) newsletter with:
Subject line theme: {req.subject_line or req.topic}
Key highlights to include: {req.highlights}

Format:
- Subject line (compelling, <60 chars)
- Preview text (<100 chars)
- Email body: greeting, intro hook, highlights (bullet points), main CTA, sign-off
- Tone: Professional but approachable, engineer-friendly
- Length: 200-300 words body

Write ONLY the email content in plain text format.""",

        "landing": f"""You are a product copywriter for Seeed Studio.

Write landing page copy for: {req.product or 'our edge AI product'}
Key features: {req.features}
Target audience: {req.audience}

Include:
- Hero headline (benefit-focused, not feature-focused)
- Subheadline
- 3-4 key benefits with short descriptions
- Social proof placeholder
- CTA section

Tone: Technical but accessible, confidence-inspiring
Write ONLY the copy, structured clearly.""",

        "blog": f"""You are a technical content writer for Seeed Studio.

Write a {req.blog_length} blog post outline + introduction for: {req.topic}

For 'short': ~800 words plan + 200-word intro
For 'medium': ~1500 words plan + 300-word intro
For 'long': ~3000 words plan + 400-word intro

Include:
- SEO-optimized title
- Meta description (155 chars)
- Full outline with H2/H3 headings
- Introduction paragraph
- Key takeaways section

Write the full structure and intro, use markdown formatting.""",

        "seo": f"""You are an SEO content strategist for Seeed Studio.

Write an SEO-optimized article for:
Primary keyword: {req.keyword}
Target region: {req.region}

Include:
- Title tag (≤60 chars, keyword-first)
- Meta description (≤155 chars)
- Article structure with H1, H2, H3 headings
- Introduction optimized for featured snippets
- FAQ section with 3-5 questions
- Natural keyword integration throughout
- Internal linking placeholders

Length: ~1200 words. Use markdown formatting.""",

        "case": f"""You are a B2B case study writer for Seeed Studio.

Write a customer success story for:
Customer: {req.customer}
Industry: {req.industry}
Challenge/Pain point: {req.pain_point}
Solution deployed: {req.solution}

Structure:
1. Executive Summary (2-3 sentences)
2. The Challenge (detailed pain points)
3. Solution Overview (how Seeed's product solved it)
4. Implementation (phases/timeline)
5. Results (use realistic metrics: X% improvement, Y% reduction)
6. Customer Quote (realistic quote)
7. Next Steps / Expansion Plans

Tone: Professional, results-focused, credible
Write the complete case study.""",

        "wechat": f"""你是微信公众号的专业科技内容编辑，为Seeed Studio（矽递科技）撰稿。

话题：{req.topic or req.product}
文章角度：{req.angle}（product=新品发布, insight=行业洞察）

要求：
- 标题：吸引眼球，可以用数字/疑问/对比，不超过20字
- 摘要：50字以内，吸引用户点开
- 正文：800-1200字，分段清晰，适当使用小标题
- 语言：专业但不枯燥，对工程师友好
- 结尾：引导互动（评论/关注/分享）
- 在合适位置使用1-2个相关Emoji

只写文章内容，用Markdown格式，中文撰写。"""
    }
    base_prompt = prompts.get(req.content_type, f"Write content about: {req.topic}")
    return base_prompt


def build_final_prompt(req: GenerateRequest, db: sqlite3.Connection) -> str:
    base_prompt = build_prompt(req)
    parts: List[str] = []

    materials_ctx = get_materials_context(req.material_ids or [], db)
    if materials_ctx:
        parts.append(
            "Use the following reference materials to keep the output accurate and specific.\n\n"
            "=== REFERENCE MATERIALS ===\n"
            f"{materials_ctx}\n"
            "=== END REFERENCE MATERIALS ==="
        )

    if (req.custom_prompt or "").strip():
        parts.append(
            "Additional writing instructions from the user:\n"
            f"{req.custom_prompt.strip()}"
        )

    parts.append(base_prompt)
    return "\n\n".join(parts)


def create_openai_client(req: GenerateRequest):
    api_key = (
        CODEX_CONFIG.get("api_key")
        or os.getenv("OPENAI_API_KEY", "")
        or req.api_key
        or ""
    )
    base_url = CODEX_CONFIG.get("base_url") or os.getenv("OPENAI_BASE_URL") or None
    model = CODEX_CONFIG.get("model") or "gpt-4o"

    if not api_key:
        raise HTTPException(status_code=400, detail="No API key found. Check ~/.codex/auth.json or .env")

    http_client = httpx.Client(proxy=None)
    client = OpenAI(api_key=api_key, base_url=base_url, http_client=http_client)
    return client, model, base_url


def create_default_openai_client():
    api_key = CODEX_CONFIG.get("api_key") or os.getenv("OPENAI_API_KEY", "") or ""
    base_url = CODEX_CONFIG.get("base_url") or os.getenv("OPENAI_BASE_URL") or None
    model = CODEX_CONFIG.get("model") or "gpt-4o"
    if not api_key:
        raise HTTPException(status_code=400, detail="No API key found. Check ~/.codex/auth.json or .env")
    http_client = httpx.Client(proxy=None)
    client = OpenAI(api_key=api_key, base_url=base_url, http_client=http_client)
    return client, model, base_url


def strip_markdown_fence(text: str) -> str:
    cleaned = (text or "").strip()
    if cleaned.startswith("```"):
        cleaned = re.sub(r"^```[a-zA-Z0-9_-]*\n?", "", cleaned)
        cleaned = re.sub(r"\n?```$", "", cleaned)
    return cleaned.strip()


def fetch_page_text(url: str) -> Dict[str, str]:
    if not HAS_BS4:
        raise HTTPException(status_code=500, detail="beautifulsoup4 not installed")
    http = httpx.Client(proxy=None, timeout=20, follow_redirects=True)
    resp = http.get(url, headers={"User-Agent": "Mozilla/5.0 (compatible; SNSBot/1.0)"})
    resp.raise_for_status()
    soup = BeautifulSoup(resp.text, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header", "aside", "noscript"]):
        tag.decompose()
    title = soup.find("title")
    title_text = title.get_text(strip=True) if title else ""
    main = soup.find("main") or soup.find("article") or soup.find("body") or soup
    text = re.sub(r"\n{3,}", "\n\n", main.get_text(separator="\n", strip=True))
    return {"title": title_text, "text": text[:18000]}


def extract_structured_json(kind: str, source: sqlite3.Row, page_title: str, page_text: str) -> Dict[str, Any]:
    client, model, _ = create_default_openai_client()
    base_context = (
        f"Source name: {source['name']}\n"
        f"Kind: {source['kind']}\n"
        f"Platform: {source['platform']}\n"
        f"Region: {source['region']}\n"
        f"URL: {source['url']}\n"
        f"Page title: {page_title}\n\n"
        f"Page text:\n{page_text}"
    )

    if kind == "monthly":
        schema_hint = """Return JSON only:
{
  "month": "April",
  "year": 2026,
  "linkedin_impressions": 0,
  "x_impressions": 0,
  "linkedin_engagement": 0,
  "x_engagement": 0,
  "linkedin_followers": 0,
  "x_followers": 0,
  "linkedin_clicks": 0,
  "x_clicks": 0,
  "raw_excerpt": "short explanation"
}
Use 0 when a metric is unavailable."""
    elif kind == "post":
        schema_hint = """Return JSON only:
{
  "posts": [
    {
      "platform": "LinkedIn",
      "title": "",
      "content": "",
      "post_type": "",
      "impressions": 0,
      "likes": 0,
      "comments": 0,
      "shares": 0,
      "engagement_rate": 0,
      "posted_at": "2026-04-16",
      "month": 4,
      "year": 2026
    }
  ]
}
Extract up to 10 recent posts visible on the page. Use 0 if a metric is missing."""
    else:
        schema_hint = """Return JSON only:
{
  "name": "",
  "platform": "",
  "region": "",
  "followers": "",
  "posting_freq": "",
  "avg_engagement": "",
  "top_actions": ["", ""],
  "learnings": "",
  "raw_excerpt": ""
}
Summarize concrete content actions from the crawled page. top_actions should contain 3-5 actionable bullets."""

    prompt = (
        "You are a data extraction engine. Read the following crawled webpage content and "
        "return valid JSON only, no markdown, no explanation.\n\n"
        f"{schema_hint}\n\n"
        f"{base_context}"
    )
    resp = client.chat.completions.create(
        model=model,
        messages=[{"role": "user", "content": prompt}],
        max_tokens=2200,
        temperature=0.1,
    )
    content = strip_markdown_fence(resp.choices[0].message.content or "")
    try:
        return json.loads(content)
    except json.JSONDecodeError as e:
        raise HTTPException(status_code=500, detail=f"Failed to parse model JSON for {kind}: {e}")


def refresh_source(source: sqlite3.Row, db: sqlite3.Connection) -> Dict[str, Any]:
    crawled = fetch_page_text(source["url"])
    data = extract_structured_json(source["kind"], source, crawled["title"], crawled["text"])

    if source["kind"] == "monthly":
        db.execute(
            """
            INSERT INTO live_monthly_snapshots (
                source_id, month, year, linkedin_impressions, x_impressions,
                linkedin_engagement, x_engagement, linkedin_followers, x_followers,
                linkedin_clicks, x_clicks, raw_excerpt
            ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)
            """,
            (
                source["id"],
                data.get("month", ""),
                int(data.get("year", datetime.now().year)),
                int(data.get("linkedin_impressions", 0) or 0),
                int(data.get("x_impressions", 0) or 0),
                float(data.get("linkedin_engagement", 0) or 0),
                float(data.get("x_engagement", 0) or 0),
                int(data.get("linkedin_followers", 0) or 0),
                int(data.get("x_followers", 0) or 0),
                int(data.get("linkedin_clicks", 0) or 0),
                int(data.get("x_clicks", 0) or 0),
                data.get("raw_excerpt", "")[:1200],
            ),
        )
        db.commit()
        return {"kind": "monthly", "source_id": source["id"], "month": data.get("month"), "year": data.get("year")}

    if source["kind"] == "post":
        inserted = 0
        for post in data.get("posts", [])[:10]:
            title = (post.get("title") or "").strip()
            if not title:
                continue
            db.execute(
                """
                INSERT INTO live_posts (
                    source_id, platform, title, content, post_type, impressions, likes,
                    comments, shares, engagement_rate, posted_at, month, year, source_url
                ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                """,
                (
                    source["id"],
                    post.get("platform") or source["platform"] or "Unknown",
                    title,
                    post.get("content", ""),
                    post.get("post_type", ""),
                    int(post.get("impressions", 0) or 0),
                    int(post.get("likes", 0) or 0),
                    int(post.get("comments", 0) or 0),
                    int(post.get("shares", 0) or 0),
                    float(post.get("engagement_rate", 0) or 0),
                    post.get("posted_at", ""),
                    int(post.get("month", 0) or 0) or None,
                    int(post.get("year", 0) or 0) or None,
                    source["url"],
                ),
            )
            inserted += 1
        db.commit()
        return {"kind": "post", "source_id": source["id"], "inserted_posts": inserted}

    db.execute(
        """
        INSERT INTO competitor_snapshots (
            source_id, name, platform, region, followers, posting_freq,
            avg_engagement, top_actions_json, learnings, raw_excerpt
        ) VALUES (?,?,?,?,?,?,?,?,?,?)
        """,
        (
            source["id"],
            data.get("name") or source["name"],
            data.get("platform") or source["platform"],
            data.get("region") or source["region"],
            data.get("followers", ""),
            data.get("posting_freq", ""),
            data.get("avg_engagement", ""),
            json.dumps(data.get("top_actions", []), ensure_ascii=False),
            data.get("learnings", ""),
            data.get("raw_excerpt", "")[:1200],
        ),
    )
    db.commit()
    return {"kind": "competitor", "source_id": source["id"], "name": data.get("name") or source["name"]}

def get_materials_context(material_ids: List[int], db: sqlite3.Connection) -> str:
    """Fetch materials from DB and build context string"""
    if not material_ids:
        return ""
    placeholders = ",".join("?" * len(material_ids))
    rows = db.execute(
        f"SELECT name, type, content FROM materials WHERE id IN ({placeholders})",
        material_ids
    ).fetchall()
    if not rows:
        return ""
    parts = []
    for r in rows:
        parts.append(f"[{r['type'].upper()}] {r['name']}\n{r['content'] or '(no content)'}")
    return "\n\n---\n\n".join(parts)

@app.post("/api/generate")
async def generate_content(req: GenerateRequest, db: sqlite3.Connection = Depends(get_db)):
    client, model, base_url = create_openai_client(req)
    prompt = build_final_prompt(req, db)

    print(f"[generate] model={model}, base_url={base_url}, content_type={req.content_type}")

    def stream_response():
        try:
            stream = client.chat.completions.create(
                model=model,
                messages=[{"role": "user", "content": prompt}],
                stream=True,
                max_tokens=2000,
                temperature=0.7,
            )
            for chunk in stream:
                if chunk.choices[0].delta.content:
                    text = chunk.choices[0].delta.content
                    yield f"data: {json.dumps({'text': text})}\n\n"
            yield "data: [DONE]\n\n"
        except Exception as e:
            yield f"data: {json.dumps({'error': str(e)})}\n\n"

    return StreamingResponse(
        stream_response(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"}
    )


@app.post("/api/generate-batch")
async def generate_batch(req: GenerateRequest, db: sqlite3.Connection = Depends(get_db)):
    platforms = [p for p in (req.platforms or ([req.platform] if req.platform else [])) if p]
    if not platforms:
        raise HTTPException(status_code=400, detail="Please select at least one platform")

    client, model, base_url = create_openai_client(req)
    results = []

    for platform in platforms:
        single_req = req.model_copy(update={"platform": platform})
        prompt = build_final_prompt(single_req, db)
        try:
            resp = client.chat.completions.create(
                model=model,
                messages=[{"role": "user", "content": prompt}],
                max_tokens=2400,
                temperature=0.7,
            )
            content = resp.choices[0].message.content or ""
            results.append({"platform": platform, "content": content})
        except Exception as e:
            results.append({"platform": platform, "error": str(e)})

    return {
        "success": True,
        "model": model,
        "base_url": base_url,
        "results": results,
    }

# ── SNS Monthly Data CRUD ──────────────────────────────────────
@app.get("/api/sns-data")
def get_sns_data(db: sqlite3.Connection = Depends(get_db)):
    rows = db.execute("SELECT * FROM sns_monthly ORDER BY year, CASE month "
                      "WHEN 'January' THEN 1 WHEN 'February' THEN 2 WHEN 'March' THEN 3 "
                      "WHEN 'April' THEN 4 WHEN 'May' THEN 5 WHEN 'June' THEN 6 "
                      "WHEN 'July' THEN 7 WHEN 'August' THEN 8 WHEN 'September' THEN 9 "
                      "WHEN 'October' THEN 10 WHEN 'November' THEN 11 WHEN 'December' THEN 12 END").fetchall()
    return [dict(r) for r in rows]

@app.post("/api/sns-data")
def upsert_sns_data(data: SNSMonthlyData, db: sqlite3.Connection = Depends(get_db)):
    db.execute("""
        INSERT INTO sns_monthly (month, year, linkedin_impressions, x_impressions,
            linkedin_engagement, x_engagement, linkedin_followers, x_followers,
            linkedin_clicks, x_clicks)
        VALUES (?,?,?,?,?,?,?,?,?,?)
        ON CONFLICT(month, year) DO UPDATE SET
            linkedin_impressions=excluded.linkedin_impressions,
            x_impressions=excluded.x_impressions,
            linkedin_engagement=excluded.linkedin_engagement,
            x_engagement=excluded.x_engagement,
            linkedin_followers=excluded.linkedin_followers,
            x_followers=excluded.x_followers,
            linkedin_clicks=excluded.linkedin_clicks,
            x_clicks=excluded.x_clicks
    """, (data.month, data.year, data.linkedin_impressions, data.x_impressions,
          data.linkedin_engagement, data.x_engagement, data.linkedin_followers,
          data.x_followers, data.linkedin_clicks, data.x_clicks))
    db.commit()
    return {"success": True, "message": f"Data for {data.month} {data.year} saved"}

@app.delete("/api/sns-data/{data_id}")
def delete_sns_data(data_id: int, db: sqlite3.Connection = Depends(get_db)):
    db.execute("DELETE FROM sns_monthly WHERE id=?", (data_id,))
    db.commit()
    return {"success": True}

# ── Materials CRUD ─────────────────────────────────────────────
@app.get("/api/materials")
def list_materials(db: sqlite3.Connection = Depends(get_db)):
    rows = db.execute(
        "SELECT id, name, type, substr(content,1,300) as preview, file_path, created_at FROM materials ORDER BY created_at DESC"
    ).fetchall()
    return [dict(r) for r in rows]

@app.delete("/api/materials/{material_id}")
def delete_material(material_id: int, db: sqlite3.Connection = Depends(get_db)):
    row = db.execute("SELECT file_path FROM materials WHERE id=?", (material_id,)).fetchone()
    if row and row["file_path"]:
        try:
            Path(row["file_path"]).unlink(missing_ok=True)
        except Exception:
            pass
    db.execute("DELETE FROM materials WHERE id=?", (material_id,))
    db.commit()
    return {"success": True}

@app.post("/api/materials/upload")
async def upload_material(
    file: UploadFile = File(...),
    db: sqlite3.Connection = Depends(get_db)
):
    filename = file.filename or "unnamed"
    ext = Path(filename).suffix.lower()
    content = ""
    file_path_str = ""

    if ext in (".txt", ".md"):
        raw = await file.read()
        content = raw.decode("utf-8", errors="replace")
        mat_type = "document"

    elif ext == ".docx":
        if not HAS_DOCX:
            raise HTTPException(status_code=500, detail="python-docx not installed")
        tmp_path = UPLOAD_DIR / filename
        with open(tmp_path, "wb") as f:
            shutil.copyfileobj(file.file, f)
        doc = DocxDocument(str(tmp_path))
        content = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        file_path_str = str(tmp_path)
        mat_type = "document"

    elif ext in (".jpg", ".jpeg", ".png", ".webp", ".gif"):
        raw = await file.read()
        tmp_path = UPLOAD_DIR / filename
        with open(tmp_path, "wb") as f:
            f.write(raw)
        file_path_str = str(tmp_path)
        mat_type = "image"

        # Vision extraction
        try:
            api_key = CODEX_CONFIG.get("api_key") or os.getenv("OPENAI_API_KEY", "")
            base_url = CODEX_CONFIG.get("base_url")
            model = CODEX_CONFIG.get("model") or "gpt-4o"
            http_client = httpx.Client(proxy=None)
            client = OpenAI(api_key=api_key, base_url=base_url, http_client=http_client)
            b64 = base64.b64encode(raw).decode()
            mime = "image/jpeg" if ext in (".jpg", ".jpeg") else f"image/{ext[1:]}"
            resp = client.chat.completions.create(
                model=model,
                messages=[{
                    "role": "user",
                    "content": [
                        {"type": "text", "text": (
                            "This is a product image for Seeed Studio. "
                            "Extract ALL useful information: product name, model number, key specs/features, "
                            "use cases, taglines, any text visible in the image. "
                            "Format as structured text for use as marketing reference material."
                        )},
                        {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{b64}"}},
                    ]
                }],
                max_tokens=1000,
            )
            content = resp.choices[0].message.content or ""
        except Exception as e:
            content = f"[Image uploaded. Vision extraction failed: {e}]"

    else:
        raise HTTPException(status_code=400, detail=f"Unsupported file type: {ext}")

    cursor = db.execute(
        "INSERT INTO materials (name, type, content, file_path) VALUES (?,?,?,?)",
        (filename, mat_type, content, file_path_str)
    )
    db.commit()
    return {
        "success": True,
        "id": cursor.lastrowid,
        "name": filename,
        "type": mat_type,
        "preview": content[:200],
    }

@app.post("/api/materials/url")
async def ingest_url(
    req: URLIngestRequest,
    db: sqlite3.Connection = Depends(get_db)
):
    if not HAS_BS4:
        raise HTTPException(status_code=500, detail="beautifulsoup4 not installed")
    try:
        http = httpx.Client(proxy=None, timeout=15, follow_redirects=True)
        resp = http.get(req.url, headers={"User-Agent": "Mozilla/5.0 (compatible; SNSBot/1.0)"})
        resp.raise_for_status()
        soup = BeautifulSoup(resp.text, "html.parser")
        # Remove nav/footer/script/style
        for tag in soup(["script", "style", "nav", "footer", "header", "aside", "noscript"]):
            tag.decompose()
        # Extract title
        title = soup.find("title")
        title_text = title.get_text(strip=True) if title else ""
        # Extract main body text
        main = soup.find("main") or soup.find("article") or soup.find("body") or soup
        text = re.sub(r'\n{3,}', '\n\n', main.get_text(separator="\n", strip=True))
        content = (f"Title: {title_text}\nURL: {req.url}\n\n" + text)[:8000]
        name = req.name or title_text or req.url[:60]
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to fetch URL: {e}")

    cursor = db.execute(
        "INSERT INTO materials (name, type, content, file_path) VALUES (?,?,?,?)",
        (name, "webpage", content, req.url)
    )
    db.commit()
    return {
        "success": True,
        "id": cursor.lastrowid,
        "name": name,
        "type": "webpage",
        "preview": content[:200],
    }


@app.get("/api/crawl-sources")
def list_crawl_sources(kind: Optional[str] = None, db: sqlite3.Connection = Depends(get_db)):
    if kind:
        rows = db.execute(
            "SELECT * FROM crawl_sources WHERE kind=? ORDER BY created_at DESC, id DESC",
            (kind,),
        ).fetchall()
    else:
        rows = db.execute("SELECT * FROM crawl_sources ORDER BY created_at DESC, id DESC").fetchall()
    return [dict(r) for r in rows]


@app.post("/api/crawl-sources")
def create_crawl_source(source: CrawlSourceIn, db: sqlite3.Connection = Depends(get_db)):
    if source.kind not in {"monthly", "post", "competitor"}:
        raise HTTPException(status_code=400, detail="kind must be monthly, post, or competitor")
    cursor = db.execute(
        """
        INSERT INTO crawl_sources (kind, name, platform, region, url)
        VALUES (?,?,?,?,?)
        ON CONFLICT(url) DO UPDATE SET
            kind=excluded.kind,
            name=excluded.name,
            platform=excluded.platform,
            region=excluded.region
        """,
        (source.kind, source.name, source.platform or "", source.region or "", source.url),
    )
    db.commit()
    row_id = cursor.lastrowid
    if not row_id:
        row = db.execute("SELECT id FROM crawl_sources WHERE url=?", (source.url,)).fetchone()
        row_id = row["id"] if row else None
    return {"success": True, "id": row_id}


@app.delete("/api/crawl-sources/{source_id}")
def delete_crawl_source(source_id: int, db: sqlite3.Connection = Depends(get_db)):
    db.execute("DELETE FROM crawl_sources WHERE id=?", (source_id,))
    db.commit()
    return {"success": True}


@app.post("/api/live-refresh")
def live_refresh(req: CrawlRefreshRequest, db: sqlite3.Connection = Depends(get_db)):
    sql = "SELECT * FROM crawl_sources"
    params: List[Any] = []
    clauses: List[str] = []
    if req.kind:
        clauses.append("kind=?")
        params.append(req.kind)
    if req.source_ids:
        placeholders = ",".join("?" * len(req.source_ids))
        clauses.append(f"id IN ({placeholders})")
        params.extend(req.source_ids)
    if clauses:
        sql += " WHERE " + " AND ".join(clauses)
    sql += " ORDER BY id DESC"

    rows = db.execute(sql, params).fetchall()
    if not rows:
        raise HTTPException(status_code=400, detail="No crawl sources matched the refresh request")

    results = []
    for row in rows:
        try:
            results.append({"success": True, **refresh_source(row, db)})
        except Exception as e:
            results.append({"success": False, "source_id": row["id"], "kind": row["kind"], "error": str(e)})
    return {"success": True, "results": results}


@app.get("/api/live-dashboard")
def get_live_dashboard(db: sqlite3.Connection = Depends(get_db)):
    monthly_rows = db.execute(
        """
        SELECT lm.*
        FROM live_monthly_snapshots lm
        JOIN (
            SELECT month, year, MAX(id) AS max_id
            FROM live_monthly_snapshots
            GROUP BY month, year
        ) latest ON latest.max_id = lm.id
        ORDER BY lm.year, CASE lm.month
            WHEN 'January' THEN 1 WHEN 'February' THEN 2 WHEN 'March' THEN 3
            WHEN 'April' THEN 4 WHEN 'May' THEN 5 WHEN 'June' THEN 6
            WHEN 'July' THEN 7 WHEN 'August' THEN 8 WHEN 'September' THEN 9
            WHEN 'October' THEN 10 WHEN 'November' THEN 11 WHEN 'December' THEN 12
            ELSE 99 END
        """
    ).fetchall()

    post_rows = db.execute(
        """
        SELECT *
        FROM live_posts
        ORDER BY
            CASE WHEN posted_at='' THEN created_at ELSE posted_at END DESC,
            id DESC
        LIMIT 30
        """
    ).fetchall()

    competitor_rows = db.execute(
        """
        SELECT cs.id AS source_id, cs.name AS source_name, cs.platform AS source_platform, cs.region AS source_region,
               cp.*
        FROM crawl_sources cs
        LEFT JOIN competitor_snapshots cp
            ON cp.id = (
                SELECT id FROM competitor_snapshots
                WHERE source_id = cs.id
                ORDER BY id DESC
                LIMIT 1
            )
        WHERE cs.kind='competitor'
        ORDER BY cs.id DESC
        """
    ).fetchall()

    competitors = []
    for row in competitor_rows:
        item = dict(row)
        raw_actions = item.get("top_actions_json") or "[]"
        try:
            item["top_actions"] = json.loads(raw_actions)
        except Exception:
            item["top_actions"] = []
        competitors.append(item)

    return {
        "monthly": [dict(r) for r in monthly_rows],
        "posts": [dict(r) for r in post_rows],
        "competitors": competitors,
    }

# ── Posts CRUD ─────────────────────────────────────────────────
@app.get("/api/posts")
def get_posts(month: Optional[int] = None, year: Optional[int] = None,
              db: sqlite3.Connection = Depends(get_db)):
    if month and year:
        rows = db.execute("SELECT * FROM posts WHERE month=? AND year=? ORDER BY impressions DESC",
                          (month, year)).fetchall()
    else:
        rows = db.execute("SELECT * FROM posts ORDER BY posted_at DESC, impressions DESC").fetchall()
    return [dict(r) for r in rows]

@app.post("/api/posts")
def create_post(post: PostData, db: sqlite3.Connection = Depends(get_db)):
    cursor = db.execute("""
        INSERT INTO posts (platform, title, content, post_type, impressions, likes,
            comments, shares, engagement_rate, posted_at, month, year)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?)
    """, (post.platform, post.title, post.content, post.post_type, post.impressions,
          post.likes, post.comments, post.shares, post.engagement_rate, post.posted_at,
          post.month, post.year))
    db.commit()
    return {"success": True, "id": cursor.lastrowid}

@app.put("/api/posts/{post_id}")
def update_post(post_id: int, post: PostData, db: sqlite3.Connection = Depends(get_db)):
    db.execute("""
        UPDATE posts SET platform=?, title=?, content=?, post_type=?, impressions=?,
            likes=?, comments=?, shares=?, engagement_rate=?, posted_at=?, month=?, year=?
        WHERE id=?
    """, (post.platform, post.title, post.content, post.post_type, post.impressions,
          post.likes, post.comments, post.shares, post.engagement_rate, post.posted_at,
          post.month, post.year, post_id))
    db.commit()
    return {"success": True}

@app.delete("/api/posts/{post_id}")
def delete_post(post_id: int, db: sqlite3.Connection = Depends(get_db)):
    db.execute("DELETE FROM posts WHERE id=?", (post_id,))
    db.commit()
    return {"success": True}

# ── KOL CRUD ───────────────────────────────────────────────────
@app.get("/api/kols")
def get_kols(db: sqlite3.Connection = Depends(get_db)):
    rows = db.execute("SELECT * FROM kols ORDER BY match_score DESC").fetchall()
    return [dict(r) for r in rows]

@app.post("/api/kols")
def create_kol(kol: KOLData, db: sqlite3.Connection = Depends(get_db)):
    cursor = db.execute("""
        INSERT INTO kols (name, handle, platform, region, category, followers,
            avg_views, engagement_rate, email, match_score, status, language, notes)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)
    """, (kol.name, kol.handle, kol.platform, kol.region, kol.category, kol.followers,
          kol.avg_views, kol.engagement_rate, kol.email, kol.match_score, kol.status,
          kol.language, kol.notes))
    db.commit()
    return {"success": True, "id": cursor.lastrowid}

@app.patch("/api/kols/{kol_id}/status")
def update_kol_status(kol_id: int, update: KOLStatusUpdate, db: sqlite3.Connection = Depends(get_db)):
    db.execute("UPDATE kols SET status=? WHERE id=?", (update.status, kol_id))
    db.commit()
    return {"success": True}

@app.delete("/api/kols/{kol_id}")
def delete_kol(kol_id: int, db: sqlite3.Connection = Depends(get_db)):
    db.execute("DELETE FROM kols WHERE id=?", (kol_id,))
    db.commit()
    return {"success": True}

@app.get("/health")
def health():
    return {
        "status": "ok",
        "version": "1.0.0",
        "codex_base_url": CODEX_CONFIG.get("base_url"),
        "codex_model": CODEX_CONFIG.get("model"),
        "codex_key_loaded": bool(CODEX_CONFIG.get("api_key")),
    }

@app.get("/api/config")
def get_ai_config():
    """Return current AI provider config for frontend display"""
    return {
        "base_url": CODEX_CONFIG.get("base_url") or os.getenv("OPENAI_BASE_URL", "(default OpenAI)"),
        "model": CODEX_CONFIG.get("model", "gpt-4o"),
        "key_source": "codex (~/.codex/auth.json)" if CODEX_CONFIG.get("api_key") else ".env / request",
        "key_loaded": bool(CODEX_CONFIG.get("api_key") or os.getenv("OPENAI_API_KEY")),
    }


# ── Web Search (DuckDuckGo HTML) ─────────────────────────────────
def search_web(query: str, num_results: int = 8) -> List[Dict[str, str]]:
    """Simple web search using DuckDuckGo HTML (no API key required)."""
    if not HAS_BS4:
        return []
    try:
        import urllib.parse
        encoded_q = urllib.parse.quote(query)
        url = f"https://duckduckgo.com/html/?q={encoded_q}&kl=wt-wt"
        http = httpx.Client(proxy=None, timeout=15, follow_redirects=True)
        headers = {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/120.0.0.0 Safari/537.36"
            )
        }
        resp = http.get(url, headers=headers)
        resp.raise_for_status()
        soup = BeautifulSoup(resp.text, "html.parser")
        results: List[Dict[str, str]] = []
        for result in soup.select(".result")[:num_results]:
            a_tag = result.select_one(".result__a")
            snippet_tag = result.select_one(".result__snippet")
            if not a_tag:
                continue
            title = a_tag.get_text(strip=True)
            link = a_tag.get("href", "")
            snippet = snippet_tag.get_text(strip=True) if snippet_tag else ""
            results.append({"title": title, "url": link, "snippet": snippet})
        return results
    except Exception as e:
        print(f"[search] web search failed for '{query}': {e}")
        return []


def fetch_article_content(url: str) -> str:
    """Fetch a single article page and return cleaned text (max 8k chars)."""
    if not HAS_BS4:
        return ""
    try:
        http = httpx.Client(proxy=None, timeout=12, follow_redirects=True)
        resp = http.get(
            url,
            headers={
                "User-Agent": (
                    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                    "AppleWebKit/537.36 (KHTML, like Gecko) "
                    "Chrome/120.0.0.0 Safari/537.36"
                )
            },
        )
        resp.raise_for_status()
        soup = BeautifulSoup(resp.text, "html.parser")
        for tag in soup(["script", "style", "nav", "footer", "header", "aside", "noscript"]):
            tag.decompose()
        main = soup.find("main") or soup.find("article") or soup.find("body") or soup
        text = re.sub(r"\n{3,}", "\n\n", main.get_text(separator="\n", strip=True))
        return text[:8000]
    except Exception as e:
        print(f"[fetch_article] failed to fetch {url}: {e}")
        return ""


# ── Market News CRUD ──────────────────────────────────────────────
@app.get("/api/market-news")
def get_market_news(
    keyword: Optional[str] = None,
    db: sqlite3.Connection = Depends(get_db),
):
    sql = "SELECT * FROM market_news_cache"
    params: List[Any] = []
    if keyword:
        sql += " WHERE keyword LIKE ?"
        params.append(f"%{keyword}%")
    sql += " ORDER BY fetched_at DESC LIMIT 200"
    rows = db.execute(sql, params).fetchall()
    return [dict(r) for r in rows]


@app.delete("/api/market-news/{news_id}")
def delete_market_news(news_id: int, db: sqlite3.Connection = Depends(get_db)):
    db.execute("DELETE FROM market_news_cache WHERE id=?", (news_id,))
    db.commit()
    return {"success": True}


@app.post("/api/market-news/refresh")
def refresh_market_news(
    keywords: List[str],
    db: sqlite3.Connection = Depends(get_db),
):
    """Search the web for each keyword and cache results."""
    all_results: List[Dict[str, Any]] = []
    for kw in keywords:
        raw = search_web(kw, num_results=8)
        for item in raw:
            db.execute(
                """
                INSERT OR IGNORE INTO market_news_cache
                    (keyword, title, snippet, url, source, published_at)
                VALUES (?,?,?,?,?,?)
                """,
                (
                    kw,
                    item["title"],
                    item["snippet"],
                    item["url"],
                    item["url"].split("/")[2] if item["url"] else "",
                    "",
                ),
            )
        db.commit()
        all_results.extend(raw)
    return {"success": True, "fetched": len(all_results), "keywords": keywords}


# ── Market Reports CRUD ────────────────────────────────────────────
@app.get("/api/market-reports")
def list_market_reports(db: sqlite3.Connection = Depends(get_db)):
    rows = db.execute(
        "SELECT id, title, industry, region, language, "
        "substr(report_content,1,120) as preview, generated_by, created_at, updated_at "
        "FROM market_reports ORDER BY updated_at DESC"
    ).fetchall()
    return [dict(r) for r in rows]


@app.get("/api/market-reports/{report_id}")
def get_market_report(report_id: int, db: sqlite3.Connection = Depends(get_db)):
    row = db.execute("SELECT * FROM market_reports WHERE id=?", (report_id,)).fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="Report not found")
    return dict(row)


@app.delete("/api/market-reports/{report_id}")
def delete_market_report(report_id: int, db: sqlite3.Connection = Depends(get_db)):
    db.execute("DELETE FROM market_reports WHERE id=?", (report_id,))
    db.commit()
    return {"success": True}


# ── Market Report Generation ───────────────────────────────────────
@app.post("/api/market-reports/generate")
async def generate_market_report(
    req: MarketReportRequest,
    db: sqlite3.Connection = Depends(get_db),
):
    client, model, base_url = create_default_openai_client()

    industry = req.industry or "AI / Edge Computing Hardware"
    region = req.region or "Global"
    is_zh = (req.language or "zh-CN").lower().startswith("zh")
    report_type = req.report_type or "comprehensive"

    # 1. Search for market data
    default_kws = [
        f"{industry} market size 2025 2026",
        f"{industry} industry trends {region}",
        f"top competitors {industry}",
    ]
    keywords = req.keywords if req.keywords else default_kws

    search_results: List[Dict[str, Any]] = []
    for kw in keywords:
        results = search_web(kw, num_results=6)
        search_results.append({"keyword": kw, "results": results})

    # 2. Build context from search snippets + live news cache
    news_rows = db.execute(
        "SELECT title, snippet, url, source FROM market_news_cache "
        "WHERE keyword IN (" + ",".join("?" * len(keywords)) + ") "
        "ORDER BY fetched_at DESC LIMIT 30",
        keywords,
    ).fetchall()
    news_items = [dict(r) for r in news_rows]

    # 3. Fetch top article pages for detailed context
    detailed_context_parts: List[str] = []
    urls_to_fetch = []
    for sr in search_results:
        for item in sr["results"][:2]:
            if item["url"] and item["url"].startswith("http"):
                urls_to_fetch.append(item["url"])

    urls_to_fetch = urls_to_fetch[:5]  # limit to 5 articles
    for url in urls_to_fetch:
        content = fetch_article_content(url)
        if content:
            detailed_context_parts.append(f"[Source: {url}]\n{content[:2000]}")

    detailed_context = "\n\n---\n\n".join(detailed_context_parts)

    # 4. Build market data summary from snippets
    market_snippets = []
    for sr in search_results:
        for item in sr["results"]:
            if item["snippet"]:
                market_snippets.append(f"- {item['title']}: {item['snippet']}")
    market_data_str = "\n".join(market_snippets[:20])

    # 5. Build prompt based on language
    if is_zh:
        prompt = f"""你是一位资深行业分析师，正在为 Seeed Studio（矽递科技）生成一份市场分析报告。

## 基本信息
- 行业: {industry}
- 目标市场/地区: {region}
- 报告类型: {report_type}
- 生成时间: {datetime.now().strftime('%Y年%m月%d日')}

## 最新网络搜索结果（行业数据 & 竞品动态）
{market_data_str or '（暂无搜索结果）'}

## 近期新闻动态
{chr(10).join(f"- {n['title']}: {n['snippet']}" for n in news_items[:10]) if news_items else '（暂无缓存新闻）'}

## 详细文章内容（来自热门链接）
{detailed_context or '（无可用文章内容）'}

## 报告要求

请生成一份结构清晰的{region}市场分析报告，包含以下章节：

1. **执行摘要**（100字以内，核心结论先行）
2. **市场规模与增长**（含数据来源）
3. **行业趋势与驱动力**（3-5个核心趋势）
4. **竞争格局分析**（主要玩家、市场份额、差异化策略）
5. **市场机会与挑战**
6. **对 Seeed Studio 的战略建议**（3-5条可执行建议）
7. **数据来源**

{("报告类型为摘要版（brief），请控制在600字以内。" if report_type == "brief" else "报告类型为综合版（comprehensive），请生成详细报告，字数800-1200字。") if report_type in ("brief", "comprehensive") else "报告类型为竞品分析版（competitor），请侧重竞争格局与差异化分析，字数600-900字。" if report_type == "competitor" else ""}

格式要求：
- 使用 Markdown 格式
- 中文撰写，专业但不晦涩
- 数据引用请注明来源（可用搜索结果中的标题作为来源）
- 结尾列出参考链接（从搜索结果中提取有效URL）
- 不要解释分析过程，只输出报告正文
"""
    else:
        prompt = f"""You are a senior industry analyst generating a market analysis report for Seeed Studio.

## Basic Info
- Industry: {industry}
- Target Market / Region: {region}
- Report Type: {report_type}
- Generated: {datetime.now().strftime('%Y-%m-%d')}

## Latest Web Search Results (Market Data & Competitor Insights)
{market_data_str or '(no search results available)'}

## Recent News & Updates
{chr(10).join(f"- {n['title']}: {n['snippet']}" for n in news_items[:10]) if news_items else '(no cached news)'}

## Detailed Article Content (from top links)
{detailed_context or '(no article content available)'}

## Report Requirements

Generate a structured market analysis report with these sections:

1. **Executive Summary** (under 100 words, key conclusions first)
2. **Market Size & Growth** (with data sources)
3. **Industry Trends & Drivers** (3-5 core trends)
4. **Competitive Landscape** (key players, market share, differentiation)
5. **Market Opportunities & Challenges**
6. **Strategic Recommendations for Seeed Studio** (3-5 actionable items)
7. **Data Sources**

{("Report type is BRIEF — keep the report under 600 words." if report_type == "brief" else "Report type is COMPREHENSIVE — generate a detailed report, 800-1200 words.") if report_type in ("brief", "comprehensive") else "Report type is COMPETITOR — focus on competitive landscape and differentiation, 600-900 words." if report_type == "competitor" else ""}

Format:
- Use Markdown
- English output
- Cite data sources (use titles from search results)
- List reference links at the end
- Output only the report body, no analysis explanation
"""

    print(f"[market-report] generating: industry={industry}, region={region}, type={report_type}")

    try:
        resp = client.chat.completions.create(
            model=model,
            messages=[{"role": "user", "content": prompt}],
            max_tokens=3000,
            temperature=0.4,
        )
        report_text = resp.choices[0].message.content or ""

        # 6. Save to DB
        title_suffix = {
            "brief": "（摘要版）",
            "comprehensive": "（综合版）",
            "competitor": "（竞品分析版）",
        }.get(report_type, "")
        title = f"{region} {industry} 市场报告{title_suffix} {datetime.now().strftime('%Y.%m.%d')}"

        cursor = db.execute(
            """
            INSERT INTO market_reports
                (title, industry, region, language, report_content, market_data, trends, competitors)
            VALUES (?,?,?,?,?,?,?,?)
            """,
            (
                title,
                industry,
                region,
                req.language or "zh-CN",
                report_text,
                market_data_str[:4000],
                json.dumps(search_results, ensure_ascii=False),
                json.dumps(
                    [{"title": n["title"], "snippet": n["snippet"]} for n in news_items[:10]],
                    ensure_ascii=False,
                ),
            ),
        )
        db.commit()
        report_id = cursor.lastrowid

        return {
            "success": True,
            "id": report_id,
            "title": title,
            "industry": industry,
            "region": region,
            "report_type": report_type,
            "report_content": report_text,
            "market_data_preview": market_data_str[:500],
            "sources_count": len(search_results),
            "model": model,
            "base_url": base_url,
        }

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Report generation failed: {e}")

if __name__ == "__main__":
    import uvicorn

    host = os.getenv("HOST", "0.0.0.0")
    port = int(os.getenv("PORT", "8000"))
    reload_enabled = os.getenv("RELOAD", "true").lower() in {"1", "true", "yes", "on"}
    uvicorn.run("main:app", host=host, port=port, reload=reload_enabled)
